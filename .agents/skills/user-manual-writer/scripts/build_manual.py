#!/usr/bin/env python3
"""Build a consistently formatted end-user manual from a JSON content file.

Usage:
    python3 build_manual.py content.json output.docx
"""
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


FONT_BODY = "Calibri"
FONT_HEADING = "Calibri"
COLOR_HEADING = RGBColor(0x1F, 0x3A, 0x5F)
COLOR_MUTED = RGBColor(0x5A, 0x5A, 0x5A)
SIZE_TITLE = Pt(28)
SIZE_H1 = Pt(18)
SIZE_H2 = Pt(13)
SIZE_BODY = Pt(11)


def set_base_style(doc):
    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = SIZE_BODY
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.15


def add_field(paragraph, field_code):
    """Insert a Word field code, such as PAGE or TOC."""
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def add_header_footer(doc, app_name, manual_title):
    section = doc.sections[0]
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header_p = section.header.paragraphs[0]
    header_p.text = ""
    header_run = header_p.add_run(f"{app_name} — {manual_title}")
    header_run.font.size = Pt(9)
    header_run.font.color.rgb = COLOR_MUTED
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer_p = section.footer.paragraphs[0]
    footer_p.text = ""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.add_run("Page ")
    add_field(footer_p, "PAGE")
    footer_p.add_run(" of ")
    add_field(footer_p, "NUMPAGES")
    for run in footer_p.runs:
        run.font.size = Pt(9)
        run.font.color.rgb = COLOR_MUTED


def add_cover_page(doc, app_name, manual_title, version, date_str):
    for _ in range(6):
        doc.add_paragraph()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(app_name)
    title_run.font.size = SIZE_TITLE
    title_run.font.bold = True
    title_run.font.color.rgb = COLOR_HEADING
    title_run.font.name = FONT_HEADING

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run(manual_title)
    subtitle_run.font.size = Pt(18)
    subtitle_run.font.color.rgb = COLOR_MUTED
    subtitle_run.font.name = FONT_HEADING

    for _ in range(4):
        doc.add_paragraph()

    version_p = doc.add_paragraph()
    version_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    version_run = version_p.add_run(f"Version {version}  |  {date_str}")
    version_run.font.size = Pt(11)
    version_run.font.color.rgb = COLOR_MUTED

    doc.add_page_break()


def add_toc(doc):
    heading = doc.add_paragraph()
    heading_run = heading.add_run("Table of Contents")
    heading_run.font.size = SIZE_H1
    heading_run.font.bold = True
    heading_run.font.color.rgb = COLOR_HEADING

    toc = doc.add_paragraph()
    add_field(toc, r'TOC \\o "1-2" \\h \\z \\u')

    note = doc.add_paragraph()
    note_run = note.add_run(
        "(Right-click above and choose “Update Field” after opening in Word "
        "to populate the table of contents.)"
    )
    note_run.font.size = Pt(9)
    note_run.font.italic = True
    note_run.font.color.rgb = COLOR_MUTED
    doc.add_page_break()


def add_heading(doc, text, size, space_before=Pt(18), style_name=None):
    paragraph = doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = size
    run.font.bold = True
    run.font.color.rgb = COLOR_HEADING
    run.font.name = FONT_HEADING
    paragraph.paragraph_format.space_before = space_before
    paragraph.paragraph_format.space_after = Pt(8)
    return paragraph


def add_intro(doc, intro_text):
    add_heading(doc, "Introduction", SIZE_H1, space_before=Pt(0), style_name="Heading 1")
    for paragraph in intro_text.split("\n\n"):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    doc.add_page_break()


def add_feature_section(doc, feature):
    add_heading(doc, feature["title"], SIZE_H1, style_name="Heading 1")

    add_heading(doc, "What it's for", SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    doc.add_paragraph(feature["purpose"])

    add_heading(doc, "Where to find it", SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    doc.add_paragraph(feature["location"])

    add_heading(doc, "Steps", SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    for step in feature["steps"]:
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(step)

    tips = feature.get("tips") or []
    if tips:
        add_heading(
            doc,
            "Tips and things to watch out for",
            SIZE_H2,
            space_before=Pt(10),
            style_name="Heading 2",
        )
        for tip in tips:
            doc.add_paragraph(tip, style="List Bullet")

    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_appendix(doc, needs_clarification):
    doc.add_page_break()
    add_heading(
        doc,
        "Appendix: Items Needing Clarification",
        SIZE_H1,
        space_before=Pt(0),
        style_name="Heading 1",
    )
    if not needs_clarification:
        doc.add_paragraph(
            "None. Every feature found during the code review was confirmed and documented above."
        )
        return

    intro = doc.add_paragraph(
        "The following were found in the underlying system but could not be confirmed "
        "confidently enough to describe in the main manual. They are listed here instead "
        "of being guessed at, so nothing in this manual describes behavior that wasn't verified."
    )
    intro.italic = True
    for item in needs_clarification:
        paragraph = doc.add_paragraph(style="List Bullet")
        run = paragraph.add_run(item["item"])
        run.bold = True
        doc.add_paragraph(item["reason"])


def build(content, output_path):
    doc = Document()
    set_base_style(doc)

    app_name = content["app_name"]
    manual_title = content.get("manual_title", "User Manual")
    version = content.get("version", "1.0")
    date_str = content.get("date", "")

    add_header_footer(doc, app_name, manual_title)
    add_cover_page(doc, app_name, manual_title, version, date_str)
    add_toc(doc)
    add_intro(doc, content["introduction"])

    for feature in content["features"]:
        add_feature_section(doc, feature)

    add_appendix(doc, content.get("needs_clarification", []))
    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: python3 build_manual.py content.json output.docx")
        sys.exit(1)

    with open(sys.argv[1], encoding="utf-8") as content_file:
        content = json.load(content_file)
    build(content, sys.argv[2])
    print(f"Wrote {sys.argv[2]}")
