import importlib.util
import json
import sys
from pathlib import Path

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


BUILDER = Path(r"C:\Users\Rovick\Desktop\project\.agents\skills\user-manual-writer\scripts\build_manual.py")
spec = importlib.util.spec_from_file_location("manual_builder", BUILDER)
module = importlib.util.module_from_spec(spec)
assert spec.loader is not None
spec.loader.exec_module(module)

TOC_ENTRIES: list[str] = []
SCREENSHOT_DIR = Path(r"C:\Users\Rovick\Desktop\project\scratch\manual-visual\screenshots")
SCREENSHOTS = {
    "Your Dashboard": [
        ("01-dashboard.png", "Dashboard guide: 1 New PR1, 2 Total Requests, 3 Pending Approval, 4 Recent Requests, 5 View. These are the controls and information described in the steps below.")
    ],
    "Creating a Purchase Request (PR1)": [
        ("02-pr1-new-request.png", "New PR1 form: 1 Request Type, 2 Purpose, 3 Date Required, 4 Priority, 5 Add Item, 6 Save Draft, 7 Submit PR1.")
    ],
    "Creating a Raw Material or Service Request Directly (PR2)": [
        ("03-planning-pr2.png", "Planning PR2 form: 1 Request Type, 2 Purpose, 3 Date Required, 4 Priority, 5 Add Line, 6 Save as Draft, 7 Submit for Approval.")
    ],
    "Checking Stock and Handing Off to Purchasing (Warehouse Validation)": [
        ("04-warehouse-validation.png", "Warehouse validation queue: 1 Queue title, 2 Search, 3 Priority filter, 4 Type filter, 5 Apply. Open an available request from this queue to continue with the validation steps.")
    ],
    "Reviewing and Approving Requests": [
        ("05-approval-queue.png", "Approval queue: 1 Approval Queue, 2 Awaiting My Action, 3 Current Step, 4 Review. Select Review to open the item before choosing an approval action.")
    ],
    "Sourcing Suppliers with a Request for Quotation (RFQ)": [
        ("06-rfq.png", "Canvassing and RFQ queue: 1 Awaiting RFQ, 2 RFQ Issued, 3 Active RFQs, 4 Open RFQ. Open the relevant RFQ to invite suppliers or compare quotations.")
    ],
    "Creating and Sending a Purchase Order": [
        ("07-purchase-orders.png", "Purchase Orders screen: 1 Generate PO, 2 Total POs, 3 For Approval, 4 Status filter, 5 View. Open a PO to review, submit, print, or send it.")
    ],
    "Submitting a Quotation (Supplier)": [
        ("08-supplier-quotations.png", "Supplier quotation list: 1 Awaiting Response, 2 RFQ number, 3 Status, 4 Submit Quote. Select Submit Quote to enter and submit your quotation.")
    ],
    "Updating Delivery Status (Supplier)": [
        ("09-supplier-deliveries.png", "Supplier deliveries list: 1 Deliveries, 2 Search, 3 Status, 4 View. Open a delivery from this list to change its status.")
    ],
    "Receiving Goods and Closing the Goods Receipt (GRN)": [
        ("10-goods-receipt.png", "Goods Receipt list: 1 Goods Receipt Notes, 2 Open, 3 Status, 4 View. Open the applicable receipt to enter quantities and submit it.")
    ],
    "Approving Quality (TSQA)": [
        ("11-tsqa.png", "TSQA dashboard: 1 TSQA Dashboard, 2 View GRN QA Queue, 3 GRNs in Queue, 4 Items Pending QA, 5 Fully Approved. Select View GRN QA Queue before recording a pass or fail result.")
    ],
    "Sharing and Reviewing Wishlist Ideas": [
        ("12-wishlist.png", "Wishlist form: 1 Submit an idea, 2 Title, 3 Description, 4 Category, 5 Submit idea.")
    ],
    "Managing Users (Admin)": [
        ("13-admin-users.png", "Admin Users screen: 1 Create User, 2 Bulk Import, 3 Search, 4 Role filter, 5 View. Use the matching control for the user-management task you need."),
        ("14-admin-bulk-import.png", "Bulk Import modal: 1 Upload stage, 2 Template information, 3 Click or drag and drop to upload, 4 Cancel.")
    ],
    "Bulk Importing Users (Admin)": [
        ("14-admin-bulk-import.png", "Bulk Import modal: 1 Upload stage, 2 Template information, 3 Click or drag and drop to upload, 4 Cancel.")
    ],
    "Viewing System Roles (Admin)": [
        ("15-admin-roles.png", "System Roles screen: 1 Roles, 2 Role Name, 3 Users, 4 Created.")
    ],
    "Configuring Approval Workflows (Admin)": [
        ("16-admin-workflows.png", "Approval Workflows screen: 1 Approval Workflows, 2 Code, 3 Name, 4 Steps. Select a workflow row before reviewing or editing its steps.")
    ],
    "System Settings (Admin)": [
        ("17-admin-settings.png", "System Settings screen: 1 VAT Rate, 2 Save VAT Rate, 3 SMTP, 4 Dropdown Options.")
    ],
}


def add_static_toc(doc):
    heading = doc.add_paragraph()
    run = heading.add_run("Table of Contents")
    run.font.size = module.SIZE_H1
    run.font.bold = True
    run.font.color.rgb = module.COLOR_HEADING
    heading.paragraph_format.space_after = Pt(8)

    for title in TOC_ENTRIES:
        item = doc.add_paragraph(style="List Bullet")
        item.paragraph_format.space_after = Pt(2)
        item.add_run(title)


def add_feature_section(doc, feature):
    module.add_heading(doc, feature["title"], module.SIZE_H1, style_name="Heading 1")
    module.add_heading(doc, "What it's for", module.SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    doc.add_paragraph(feature["purpose"])
    module.add_heading(doc, "Where to find it", module.SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    doc.add_paragraph(feature["location"])
    shots = SCREENSHOTS.get(feature["title"], [])
    if shots:
        module.add_heading(doc, "Screenshot guide", module.SIZE_H2, space_before=Pt(10), style_name="Heading 2")
        for filename, _caption in shots:
            image_path = SCREENSHOT_DIR / filename
            if not image_path.exists():
                raise FileNotFoundError(f"Missing screenshot: {image_path}")
            picture = doc.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.keep_with_next = True
            picture.add_run().add_picture(str(image_path), width=Inches(6.4))
            caption_paragraph = doc.add_paragraph()
            caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption_paragraph.paragraph_format.space_after = Pt(8)
            caption_run = caption_paragraph.add_run("Screen shown in this section.")
            caption_run.italic = True
            caption_run.font.size = Pt(8.5)
    module.add_heading(doc, "Steps", module.SIZE_H2, space_before=Pt(10), style_name="Heading 2")
    for number, step in enumerate(feature["steps"], start=1):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.28)
        paragraph.paragraph_format.first_line_indent = Inches(-0.28)
        paragraph.add_run(f"{number}. ").bold = False
        paragraph.add_run(step)
    tips = feature.get("tips") or []
    if tips:
        module.add_heading(doc, "Tips and things to watch out for", module.SIZE_H2, space_before=Pt(10), style_name="Heading 2")
        for tip in tips:
            doc.add_paragraph(tip, style="List Bullet")
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_clean_manual.py content.json output.docx")
    content_path, output_path = map(Path, sys.argv[1:])
    content = json.loads(content_path.read_text(encoding="utf-8"))
    global TOC_ENTRIES
    TOC_ENTRIES = ["Introduction"] + [item["title"] for item in content["features"]] + ["Appendix: Items Needing Clarification"]
    module.add_toc = add_static_toc
    module.add_feature_section = add_feature_section
    module.build(content, output_path)
    print(f"Wrote {output_path}")


if __name__ == "__main__":
    main()
