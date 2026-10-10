import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

rows = [
    ("Airine Edianel", "airine.edianel@fortunepackaging.com"),
    ("Alex M. Terrones", "alex.terrones@fortunepackaging.com"),
    ("Analyn A. Ilagan", "analyn.ilagan@fortunepackaging.com"),
    ("Andrew Anastacio", "andrew.anastacio@fortunepackaging.com"),
    ("AP Importation", "apimportation@fortunepackaging.com"),
    ("Ar-jay Nebres", "arjay.nebres@printwell.com.ph"),
    ("Benjamin L. Yam", "benjamin.yam@fortunepackaging.com"),
    ("Bernard Avilla", "bernard.avilla@fortunepackaging.com"),
    ("Christian Dangate", "fpc_hrma@fortunepackaging.com"),
    ("Dexter Bariso", "dexter.bariso@fortunepackaging.com"),
    ("Edgardo C. Bautista", "edgardo.bautista@fortunepackaging.com"),
    ("Ella May P. Ubaldo", "ellamay.ubaldo@fortunepackaging.com"),
    ("Ericson Martinez", "ericson.martinez@fortunepackaging.com"),
    ("Euclid Mondejar", "euclid.mondejar@fortunepackaging.com"),
    ("Evelyn S. Gillego", "evelyn.gillego@fortunepackaging.com"),
    ("FPC Accounting", "fpc_accounting@fortunepackaging.com"),
    ("FPC Accounts Payable", "accountspayable@fortunepackaging.com"),
    ("FPC Accounts Receivable", "accountsreceivable@fortunepackaging.com"),
    ("FPC DCC", "fpc.dcc@fortunepackaging.com"),
    ("FPC Digital", "digital@fortunepackaging.com"),
    ("FPC Engineering", "fpc.engineering@fortunepackaging.com"),
    ("FPC ESH", "fpc_esh@fortunepackaging.com"),
    ("FPC PrePress", "fpc.prepress@fortunepackaging.com"),
    ("FPC SPC", "fpc.spc@fortunepackaging.com"),
    ("Genevieve Y. Roxas-chua", "genevieve.roxas-chua@fortunepackaging.com"),
    ("Grace C. Lazarte", "grace.lazarte@fortunepackaging.com"),
    ("Grace Magsino", "grace.magsino@fortunepackaging.com"),
    ("Howard M. Yam", "howard.yam@fortunepackaging.com"),
    ("Iris V. Goce", "iris.goce@fortunepackaging.com"),
    ("Jim Dominic Pabalate", "jimdominic.pabalate@fortunepackaging.com"),
    ("Jofelyn Paalisbo", "jofelyn.paalisbo@fortunepackaging.com"),
    ("Jonathan Balog", "cost.analyst@fortunepackaging.com"),
    ("Jr. Cost Estimator", "cost.estimator@fortunepackaging.com"),
    ("Laarni M. Rull", "laarni.rull@fortunepackaging.com"),
    ("Lalaine Diaz", "cost_estimator@fortunepackaging.com"),
    ("Leah Z. Gemora", "leah.gemora@fortunepackaging.com"),
    ("Ma. Dinah Amor", "dinah.amor@fortunepackaging.com"),
    ("Marcial Amaro VI", "marcial.amaro@fortunepackaging.com"),
    ("Marcial Amaro VI", "marcial.amaro@printwell.com.ph"),
    ("Maria Gilina Villamor", "mariagilina.villamor@fortunepackaging.com"),
    ("Mark Omambac", "fpc_training@fortunepackaging.com"),
    ("Marthy P. Mercado", "marthy.mercado@fortunepackaging.com"),
    ("Materials FGW", "materials.fgw@fortunepackaging.com"),
    ("Mavy Abog", "fpc_recruitment@fortunepackaging.com"),
    ("Micka Malacapo", "micka.malacapo@fortunepackaging.com"),
    ("Pincky Agdon", "pincky.trajano@fortunepackaging.com"),
    ("Rebecca Ann Gabion", "rebecca.gabion@fortunepackaging.com"),
    ("Renielle Anne Geoff M. Cunanan", "renielle.cunanan@fortunepackaging.com"),
    ("Romeo P. Llona Jr.", "romeo.llona@fortunepackaging.com"),
    ("Ronaldo Capinpin", "ronaldo.capinpin@fortunepackaging.com"),
    ("Rosemarie B. Sison", "rosemarie.sison@fortunepackaging.com"),
    ("Ryan Yam", "ryan.yam@fortunepackaging.com"),
    ("SGC Planning", "sgc_planning@fortunepackaging.com"),
    ("Sherlyn M. Maling", "sherlyn.maling@fortunepackaging.com"),
    ("Technical Support", "technical.support@fortunepackaging.com"),
    ("Tiburcio B. Estioko", "tiburcio.estioko@fortunepackaging.com"),
    ("Valerie Pabalinas", "valerie.pabalinas@fortunepackaging.com"),
]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Password Reset"

headers = ["#", "Full Name", "Email", "Password Reset", "Login Verified", "Date Reset"]
header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
body_font = Font(name="Arial")

for col, h in enumerate(headers, start=1):
    c = ws.cell(row=1, column=col, value=h)
    c.font = header_font
    c.fill = header_fill
    c.alignment = Alignment(horizontal="center", vertical="center")

for i, (name, email) in enumerate(rows, start=1):
    r = i + 1
    ws.cell(row=r, column=1, value=i).font = body_font
    ws.cell(row=r, column=2, value=name).font = body_font
    ws.cell(row=r, column=3, value=email).font = body_font
    ws.cell(row=r, column=4, value="Yes").font = body_font
    ws.cell(row=r, column=5, value="Yes").font = body_font
    ws.cell(row=r, column=6, value="2026-09-24").font = body_font

widths = [5, 30, 38, 20, 16, 14]
for col, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(col)].width = w

ws.freeze_panes = "A2"

note = ws.cell(row=len(rows) + 3, column=1,
                value="Note: A shared temporary password was applied via the Supabase admin API to all accounts "
                      "from the 2026-09-16 bulk-import batch that succeeded (57 of 65 rows; 8 failed due to "
                      "pre-existing accounts and are not included here). Login was verified for every row. "
                      "The password itself is intentionally not included in this file — share it with users "
                      "through a separate secure channel.")
note.font = Font(name="Arial", italic=True, size=9, color="808080")
ws.merge_cells(start_row=len(rows) + 3, start_column=1, end_row=len(rows) + 3, end_column=6)

out_path = "scratch/password-reset/Bulk-Import-Password-Reset.xlsx"
wb.save(out_path)
print("saved", out_path)
