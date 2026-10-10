"""Feature-specific illustrated layout using the retained manual skill template."""
import importlib.util
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('manual_template', ROOT / '.agents/skills/user-manual-writer/scripts/build_manual.py')
template = importlib.util.module_from_spec(spec)
spec.loader.exec_module(template)
template.COLOR_HEADING = RGBColor(0, 0, 0)
content = json.loads((WORK / 'content.json').read_text(encoding='utf-8'))

# Excerpts preserve the supplied UI content and enlarge the relevant controls.
assets = [
 ('d6fc3b33-175f-43e4-8007-11aa134a2f82', (278, 380, 970, 590)),
 ('e8627e70-ef48-48ae-892d-bd293e5a1496', (25, 82, 715, 237)),
 ('6323a1ab-b4ff-4707-9e25-75b822178dc5', (2, 15, 770, 294)),
 ('64478ec5-ce0f-4b25-9143-bcd44b45c3e2', (25, 75, 755, 270)),
]
for n, (identifier, box) in enumerate(assets, 1):
    source = Path('C:/Users/Rovick/AppData/Local/Temp') / f'codex-clipboard-{identifier}.png'
    Image.open(source).crop(box).save(WORK / f'figure-{n}.png')

doc = Document()
template.set_base_style(doc)
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.75)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
for style in doc.styles:
    props = style.element.find(qn('w:pPr'))
    if props is not None:
        for border in list(props.findall(qn('w:pBdr'))): props.remove(border)
doc.styles['Caption'].font.color.rgb = RGBColor(0,0,0)
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    doc.styles[name].font.name = 'Calibri'
    doc.styles[name].font.color.rgb = RGBColor(0,0,0)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.core_properties.title = content['manual_title']
doc.core_properties.subject = 'Warehouse staff instructions for editing existing request item codes'
doc.core_properties.author = 'Fortune Procurement System'
doc.core_properties.keywords = 'warehouse, item code, user guide'
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Warehouse Item Code User Guide  |  ')
template.add_field(footer, 'PAGE')
for run in footer.runs: run.font.size = Pt(9)

def heading(text, level=1):
    return template.add_heading(doc, text, Pt(17 if level == 1 else 11), Pt(10), f'Heading {level}')

def figure(number, caption, width=6.6):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(WORK / f'figure-{number}.png'), width=Inches(width))
    pic = p._p.xpath('.//wp:docPr')[0]
    pic.set('descr', caption)
    p = doc.add_paragraph(f'Figure {number}. {caption}', style='Caption')
    p.paragraph_format.space_after = Pt(7)
    for run in p.runs: run.font.size = Pt(9)

def feature(index, screenshots=()):
    f = content['features'][index]
    heading(f['title'])
    for label, text in [('What it is for', f['purpose']), ('Where to find it', f['location'])]:
        heading(label, 2)
        doc.add_paragraph(text)
    heading('Steps', 2)
    for n, step in enumerate(f['steps'], 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.22)
        p.paragraph_format.first_line_indent = Inches(-.22)
        p.add_run(f'{n}. {step}')
    for args in screenshots: figure(*args)
    if f['tips']:
        heading('Tips and things to watch out for', 2)
        for tip in f['tips']: doc.add_paragraph(tip, style='List Bullet')

p = doc.add_paragraph(content['manual_title'], style='Title')
for r in p.runs: r.font.size = Pt(25); r.font.bold = True; r.font.color.rgb = RGBColor(0,0,0)
doc.add_paragraph(f"{content['app_name']}  |  Version {content['version']}  |  {content['date']}")
for text in content['introduction'].split('\n\n'): doc.add_paragraph(text)
doc.add_paragraph('In this guide: Open the editor • Enter and save • Cancel • Remove a code • Reload and resolve errors')
feature(0, [(1, 'The CODE column shows (blank) and the Edit code control.', 6.5)])
doc.add_page_break()
feature(1, [(2, 'The editor opens with the current value, which is blank in this example.', 6.3),
            (3, 'Enter the intended code, then choose Save code.', 6.3)])
doc.add_page_break()
figure(4, 'A successful save displays Item code saved and the new code on the same item.', 5.8)
feature(2)
feature(3)
doc.add_page_break()
feature(4)
heading('Before you continue')
doc.add_paragraph('Confirm the saved code belongs to the correct item. Continue with the usual warehouse quantity checks, notes and validation decision. Saving an item code alone does not complete warehouse validation.')
output = ROOT / 'docs/Fortune-Warehouse-Item-Code-User-Guide.docx'
doc.save(output)
print(output)
print('Documented features: 5; clarification items: 0; supplied screenshots: 4')
